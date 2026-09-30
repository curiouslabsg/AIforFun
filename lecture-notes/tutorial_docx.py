"""Word (.docx) version of a tutorial, from the same YAML as the PDF.

Figures are SVGs, which Word cannot embed reliably, so they are rendered to PNG
first with headless Chromium (svg2png.js).
"""
import os, re, subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
LETTERS = "ABCD"
ROMAN = ["I", "II", "III", "IV"]
NAVY = RGBColor(0x0E, 0x1B, 0x33)
MUTED = RGBColor(0x5B, 0x64, 0x79)
FILL = {"cyan": "E2F5FB", "yellow": "FFF5D9", "green": "DEF6EE", "red": "FDE8E4", "purple": "EFE8FB", "deepblue": "E3ECFA"}
TOKEN = re.compile(r"(\*\*.+?\*\*|==.+?==|\^\{.+?\}|_\{.+?\}|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*]))")


def runs(par, text, size=None, bold=None, color=None):
    """Add inline-marked-up text (**bold**, *italic*, ==highlight==, x^{2}, F_{1}) to a paragraph."""
    for piece in TOKEN.split(str(text)):
        if not piece:
            continue
        r = None
        if piece.startswith("**") or piece.startswith("=="):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("^{"):
            r = par.add_run(piece[2:-1]); r.font.superscript = True
        elif piece.startswith("_{"):
            r = par.add_run(piece[2:-1]); r.font.subscript = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1]); r.italic = True
        else:
            r = par.add_run(piece)
        if size: r.font.size = Pt(size)
        if bold is not None and not r.bold: r.bold = bold
        if color is not None: r.font.color.rgb = color
    return par


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def no_borders(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil"); borders.append(el)
    tbl.tblPr.append(borders)


def para(doc, text="", size=10.5, bold=None, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    runs(p, text, size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p


def dotted_lines(doc, n, indent=Mm(16)):
    for _ in range(n):
        p = doc.add_paragraph("." * 110)
        p.paragraph_format.left_indent = indent
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs: r.font.color.rgb = MUTED; r.font.size = Pt(9)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def pngs(doc_yaml, env):
    """Render every figure the tutorial uses to PNG, once."""
    names = set()
    def walk(x):
        if isinstance(x, dict):
            for k in ("figure",):
                if isinstance(x.get(k), str): names.add(x[k])
            for v in x.values(): walk(v)
        elif isinstance(x, list):
            for v in x: walk(v)
    walk(doc_yaml)
    names.add("cover-light.svg")
    out = os.path.join(HERE, "output", "png")
    os.makedirs(out, exist_ok=True)
    todo = [os.path.join(HERE, "diagrams", n) for n in sorted(names)]
    subprocess.run(["node", os.path.join(HERE, "svg2png.js"), out] + todo, check=True, env=env)
    return {n: os.path.join(out, n.replace(".svg", ".png")) for n in names}


def picture(doc, path, width_mm):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Mm(width_mm))


def answer(doc, ans, img):
    if ans is None:
        return
    for a in ans if isinstance(ans, list) else [ans]:
        if isinstance(a, dict) and "figure" in a:
            picture(doc, img[a["figure"]], 150 * a.get("width", 100) / 100)
        elif isinstance(a, dict):
            p = para(doc, a.get("do", ""), 9.5, space_after=0)
            p.paragraph_format.left_indent = Mm(8)
            q = para(doc, a.get("work", ""), 9.5, bold=True, color=RGBColor(0x0A, 0x7D, 0x5F), space_after=3)
            q.paragraph_format.left_indent = Mm(8)
        else:
            p = para(doc, a, 9.5, space_after=3); p.paragraph_format.left_indent = Mm(8)


def total(q):
    return sum(int(p.get("marks", 0)) for p in q["parts"])


def write_docx(y, path, env):
    from tutorial import P1_INSTR, P2_INSTR
    img = pngs(y, env)
    m = y["meta"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(20), Mm(25), Mm(18), Mm(18)
    st = doc.styles["Normal"]; st.font.name = "Nunito"; st.font.size = Pt(10.5)

    # cover
    para(doc, f"{m['series']} · {m['course']}", 9, True, MUTED)
    para(doc, f"Tutorial {m['number']}", 16, True, RGBColor(0xC8, 0x8A, 0x00), space_after=0)
    para(doc, m["title"], 44, True, NAVY, space_after=2)
    para(doc, m["subtitle"], 12, color=MUTED, space_after=10)
    qt = doc.add_table(rows=1, cols=1); shade(qt.cell(0, 0), "FFF5D9")
    c = qt.cell(0, 0); c.paragraphs[0].text = ""
    runs(c.paragraphs[0], "“" + m["quote"]["text"].strip("“”‘’\"") + "”", 11.5).italic = True
    runs(c.add_paragraph(), "— " + m["quote"]["by"], 9, True, MUTED)
    para(doc, "")
    picture(doc, img["cover-light.svg"], 140)
    para(doc, "[ Infographic placeholder · drop the NotebookLM infographic here ]", 9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # concept summary: 2 × 3 grid of shaded boxes
    para(doc, "Concept summary", 18, True, NAVY, space_after=6)
    boxes = y["summary"]
    t = doc.add_table(rows=(len(boxes) + 1) // 2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, b in enumerate(boxes):
        cell = t.cell(i // 2, i % 2); shade(cell, FILL[b["colour"]])
        cell.paragraphs[0].text = ""
        runs(cell.paragraphs[0], b["title"], 11.5, True, NAVY)
        for it in b["items"]:
            runs(cell.add_paragraph(style="List Bullet"), it, 9)
    page_break(doc)

    def paper_cover(n, title, time, marks, instr, exam=None):
        para(doc, f"{m['series']} · {m['course']}", 9, True, MUTED)
        para(doc, f"Tutorial {m['number']} · {m['title']}", 11, True, MUTED, space_after=0)
        para(doc, f"Paper {n} · {title}", 24, True, NAVY, space_after=0)
        para(doc, f"{time} · {marks} marks", 11, True, RGBColor(0xEF, 0x5A, 0x44), space_after=10)
        ct = doc.add_table(rows=3, cols=2); ct.style = "Table Grid"
        for r, lab in enumerate(["Name", "Class", "Register number"]):
            ct.cell(r, 0).text = lab; ct.cell(r, 0).paragraphs[0].runs[0].bold = True
        para(doc, "")
        it = doc.add_table(rows=1, cols=1); it.style = "Table Grid"
        c = it.cell(0, 0); c.paragraphs[0].text = ""
        runs(c.paragraphs[0], "READ THESE INSTRUCTIONS FIRST", 10, True)
        for s in instr:
            runs(c.add_paragraph(style="List Bullet"), s, 10)
        if exam:
            para(doc, "")
            et = doc.add_table(rows=len(exam) + 2, cols=2); et.style = "Table Grid"
            et.cell(0, 0).text = "For examiner’s use"
            for r, (lab, mk) in enumerate(exam, 1):
                et.cell(r, 0).text = lab; et.cell(r, 1).text = f"/ {mk}"
            et.cell(len(exam) + 1, 0).text = "Total"; et.cell(len(exam) + 1, 1).text = f"/ {marks}"
        page_break(doc)

    # paper 1
    paper_cover(1, "Multiple Choice", y["paper1"]["time"], 20, P1_INSTR)
    for i, q in enumerate(y["paper1"]["questions"], 1):
        p = para(doc, "", space_after=2)
        runs(p, f"**{i}**\t{q['stem']}\t(      )")
        p.paragraph_format.tab_stops.add_tab_stop(Mm(8))
        p.paragraph_format.tab_stops.add_tab_stop(Mm(165))
        if q.get("figure"): picture(doc, img[q["figure"]], 110)
        pat, opts = q["pattern"], q.get("options", [])
        if pat == "roman":
            for k, s in enumerate(q["statements"]):
                sp = para(doc, f"**{ROMAN[k]}**\t{s}", 10, space_after=0); sp.paragraph_format.left_indent = Mm(10)
        if pat in ("grid", "roman", "row"):
            cols = 4 if pat == "row" else 2
            ot = doc.add_table(rows=len(opts) // cols, cols=cols); no_borders(ot)
            for k, o in enumerate(opts):
                cell = ot.cell(k // cols, k % cols); cell.paragraphs[0].text = ""
                runs(cell.paragraphs[0], f"**{LETTERS[k]}**  {o}", 10)
        elif pat == "list":
            for k, o in enumerate(opts):
                op = para(doc, f"**{LETTERS[k]}**  {o}", 10, space_after=0); op.paragraph_format.left_indent = Mm(10)
        elif pat == "table":
            tb = q["table"]
            tt = doc.add_table(rows=len(tb["rows"]) + 1, cols=len(tb["headers"])); tt.style = "Table Grid"
            for k, h in enumerate(tb["headers"]):
                tt.cell(0, k).paragraphs[0].text = ""; runs(tt.cell(0, k).paragraphs[0], h, 9.5, True)
            for r, row in enumerate(tb["rows"], 1):
                tt.cell(r, 0).text = LETTERS[r - 1]; tt.cell(r, 0).paragraphs[0].runs[0].bold = True
                for k, v in enumerate(row, 1):
                    tt.cell(r, k).paragraphs[0].text = ""; runs(tt.cell(r, k).paragraphs[0], v, 9.5)
        para(doc, "", space_after=4)
    para(doc, "END OF PAPER 1", 9, True, MUTED, WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # paper 2
    A, B = y["paper2"]["sectionA"], y["paper2"]["sectionB"]
    exam = [(f"Section A · Q{i}", total(q)) for i, q in enumerate(A, 1)] + [("Section B", total(B[0]))]
    paper_cover(2, "Structured Questions", y["paper2"]["time"], sum(total(q) for q in A) + total(B[0]), P2_INSTR, exam)

    def structured(q, n):
        para(doc, f"**{n}**\t{q['stem']}", space_after=3)
        if q.get("table"):
            tb = q["table"]
            tt = doc.add_table(rows=len(tb["rows"]) + 1, cols=len(tb["headers"])); tt.style = "Table Grid"
            for k, h in enumerate(tb["headers"]):
                tt.cell(0, k).text = h
            for r, row in enumerate(tb["rows"], 1):
                for k, v in enumerate(row):
                    tt.cell(r, k).text = v
        if q.get("figure"):
            picture(doc, img[q["figure"]], 150)
        for i, p in enumerate(q["parts"]):
            pp = para(doc, f"**({'abcdefgh'[i]})**\t{p['text']}\t**[{p['marks']}]**", space_after=2)
            pp.paragraph_format.left_indent = Mm(8)
            pp.paragraph_format.tab_stops.add_tab_stop(Mm(16))
            pp.paragraph_format.tab_stops.add_tab_stop(Mm(165))
            if p.get("calc"):
                dotted_lines(doc, 2)
                r = para(doc, f"{p['calc']['label']} = ........................... {p['calc'].get('unit', '')}", align=WD_ALIGN_PARAGRAPH.RIGHT)
            elif p.get("lines"):
                dotted_lines(doc, int(p["lines"]))
        para(doc, f"**[Total: {total(q)}]**", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

    para(doc, "Section A · Answer all questions.", 14, True, NAVY)
    for i, q in enumerate(A, 1):
        structured(q, i)
    page_break(doc)
    para(doc, "Section B · Answer one question.", 14, True, NAVY)
    for i, q in enumerate(B, 1):
        structured(q, len(A) + i)
    para(doc, "END OF PAPER 2", 9, True, MUTED, WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # answer key
    para(doc, "Answer key", 18, True, NAVY)
    para(doc, "Paper 1 · Multiple choice", 12, True, RGBColor(0x00, 0x74, 0x9A))
    kt = doc.add_table(rows=21, cols=3); kt.style = "Table Grid"
    for k, h in enumerate(["Q", "Answer", "Why"]):
        kt.cell(0, k).text = h; kt.cell(0, k).paragraphs[0].runs[0].bold = True; shade(kt.cell(0, k), "E2F5FB")
    for i, q in enumerate(y["paper1"]["questions"], 1):
        kt.cell(i, 0).text = str(i); kt.cell(i, 1).text = q["answer"]
        kt.cell(i, 2).paragraphs[0].text = ""; runs(kt.cell(i, 2).paragraphs[0], q["explain"], 9.5)
    for n, q in enumerate(A + B, 1):
        para(doc, "")
        para(doc, f"Paper 2 · Question {n}  [{total(q)}]", 12, True, RGBColor(0x0A, 0x7D, 0x5F))
        for i, p in enumerate(q["parts"]):
            para(doc, f"**({'abcdefgh'[i]})**  [{p['marks']}]", 10, space_after=1)
            answer(doc, p.get("answer"), img)

    doc.save(path)
    return path
